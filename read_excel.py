import openpyxl

wb = openpyxl.load_workbook(r'D:\berkas magang\SIDUL\blackbox_multisheet_sidul.xlsx')
print('Sheet names:', wb.sheetnames)

for sn in wb.sheetnames:
    ws = wb[sn]
    print(f'\n=== {sn} ===')
    print(f'Max row: {ws.max_row}, Max col: {ws.max_column}')
    
    for row in ws.iter_rows(min_row=1, max_row=min(20, ws.max_row)):
        for cell in row:
            print(f'  [{cell.row},{cell.column}] value={cell.value} | font={cell.font.name},{cell.font.size},{cell.font.bold} | fill={cell.fill.start_color.rgb if cell.fill.start_color else "none"} | align_h={cell.alignment.horizontal} | border={cell.border}')

print('\n\n=== MERGED CELLS ===')
for sn in wb.sheetnames:
    ws = wb[sn]
    print(f'{sn}: {ws.merged_cells.ranges}')
